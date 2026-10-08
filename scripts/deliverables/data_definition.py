"""데이터 정의서(엑셀) 생성 — schema.prisma와 마이그레이션 SQL이 원천이다.

    python3 scripts/deliverables/data_definition.py   → docs/deliverables/데이터정의서.xlsx

스키마가 바뀌면 다시 돌린다. 손으로 고치지 않는다(한 사실은 한 곳에만 — D38).
설명이 없는 컬럼은 아래 GLOSSARY로, 테이블 이름은 TABLES로 채운다.
"""
import re
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook

import xlsx_style as X

ROOT = Path(__file__).resolve().parents[2]
SCHEMA = ROOT / 'apps/api/prisma/schema.prisma'
MIGR = ROOT / 'apps/api/prisma/migrations'
SRC = ROOT / 'apps/api/src'
OUT = ROOT / 'docs/deliverables/데이터정의서.xlsx'

# 테이블 논리명·설명 — 스키마 주석과 architecture/entities.md 기준
TABLES = {
    'Account': ('계정', '로그인 주체. 역할 4개(본사 관리자·지점 관리자·직원·회원) 중 하나를 가진다'),
    'RefreshToken': ('리프레시 토큰', '로그인 유지용 토큰(1회용 교체). 원문 대신 해시를 저장'),
    'Branch': ('지점(위탁계약 현장)', '아파트·오피스텔과 맺은 위탁운영 계약 현장. 계약 상대방·기간·상태를 가진다'),
    'Staff': ('직원', '본사가 채용한 인력. 현재 파견 지점을 캐시로 가진다'),
    'StaffAssignment': ('파견 이력', '직원이 언제 어느 지점에 파견됐는지. 진행 중 파견은 직원당 1건'),
    'StaffProfileChangeRequest': ('인사 변경 신청', '직원 본인의 인사정보 변경 신청과 관리자 승인'),
    'AttendanceRecord': ('근태 기록', '직원별 하루 1건의 출퇴근·상태 기록'),
    'LeaveRequest': ('휴가 신청', '연차·병가 등 휴가 신청과 승인·반려'),
    'LeaveBalance': ('연차 잔여', '직원·연도별 부여 일수와 사용 일수'),
    'WorkLog': ('업무일지', '직원별 하루 업무 기록'),
    'Post': ('게시글', '본사→현장, 현장→이용자 공지'),
    'Attachment': ('첨부파일', '게시글 첨부(교육자료 등)'),
    'PostReadReceipt': ('읽음 확인', '게시글별 열람 기록'),
    'Member': ('회원', '지점에 등록된 이용자'),
    'CourseEnrollment': ('수강 등록', '회원의 프로그램 수강 내역'),
    'PTSession': ('PT 세션권', '회원이 구매한 PT 횟수와 사용 횟수'),
    'PTSessionLog': ('PT 사용 기록', 'PT 세션 1회 차감 기록'),
    'Instructor': ('강사', '지점 강사 프로필'),
    'Program': ('프로그램', '지점이 운영하는 프로그램. 이용 방식(자유이용·회차 예약·PT)과 상태를 가진다'),
    'ProgramScheduleTemplate': ('프로그램 시간표 틀', '요일별 반복 회차 틀'),
    'ScheduleSlot': ('회차', '날짜·시간·정원이 정해진 예약 단위'),
    'Reservation': ('예약', '회원의 회차 예약. 신청→확정→완료/취소/노쇼'),
    'Payment': ('결제', '예약당 1건의 모의 결제. 공급가액·부가세 분리'),
    'Facility': ('시설', '헬스장·수영장 등 지점 시설과 현재 혼잡도'),
    'CongestionSnapshot': ('혼잡도 기록', '시설별 혼잡도 시계열'),
    'FacilityCheckIn': ('시설 입퇴장', '회원의 시설 체크인·체크아웃'),
    'Asset': ('자산·비품', '지점 자산(고정자산·소모품)과 상태'),
    'Document': ('문서', '계약서·인사서류 등 문서와 보존기한'),
    'CodeSequence': ('채번 시퀀스', '지점·종류별 코드(회원번호·직원코드·자산코드) 일련번호'),
    'MemberLinkAttempt': ('회원 연동 시도', '오프라인 회원↔앱 계정 연동 실패 기록(시도 제한용)'),
    'AuditLog': ('변경 이력', '주요 엔티티 변경 기록'),
}

# 주석이 없는 컬럼의 일반 설명(이름 기준)
GLOSSARY = {
    'id': '식별자(UUID)', 'createdAt': '생성 일시', 'updatedAt': '수정 일시', 'deletedAt': '삭제 일시(소프트 삭제)',
    'accountId': '계정 ID', 'branchId': '지점 ID', 'staffId': '직원 ID', 'memberId': '회원 ID', 'programId': '프로그램 ID',
    'facilityId': '시설 ID', 'instructorId': '강사 ID', 'postId': '게시글 ID', 'reservationId': '예약 ID',
    'scheduleSlotId': '회차 ID', 'ptSessionId': 'PT 세션권 ID', 'authorId': '작성자 계정 ID', 'approverId': '승인자 ID',
    'actorId': '행위자 계정 ID', 'reviewedBy': '검토자 ID', 'entityId': '대상 레코드 ID',
    'name': '이름', 'phone': '연락처', 'address': '주소', 'birthDate': '생년월일', 'gender': '성별', 'memo': '메모', 'note': '비고',
    'status': '상태(코드 정의 참고)', 'role': '역할', 'type': '유형', 'category': '분류', 'kind': '종류', 'scope': '공개 범위',
    'isActive': '사용 여부', 'title': '제목', 'content': '내용', 'description': '설명', 'reason': '사유',
    'email': '이메일', 'passwordHash': '비밀번호 해시(bcrypt)', 'failedLoginCount': '연속 로그인 실패 횟수(미사용)',
    'lockedUntil': '로그인 잠금 해제 시각(미사용)', 'lastLoginAt': '마지막 로그인 일시', 'tokenHash': '토큰 해시',
    'expiresAt': '만료 일시', 'revokedAt': '폐기 일시',
    'contractStartAt': '계약 시작일', 'contractEndAt': '계약 종료일', 'contractStatus': '계약 상태',
    'staffCode': '직원 코드', 'hireDate': '입사일', 'resignDate': '퇴사일', 'startDate': '시작일', 'endDate': '종료일',
    'field': '변경 항목', 'oldValue': '변경 전 값', 'newValue': '변경 후 값', 'reviewedAt': '검토 일시',
    'date': '날짜', 'checkInAt': '체크인 일시', 'checkOutAt': '체크아웃 일시', 'days': '일수',
    'year': '연도', 'totalDays': '부여 일수', 'usedDays': '사용 일수',
    'pinned': '상단 고정 여부', 'viewCount': '조회수', 'fileName': '파일명', 'fileSize': '파일 크기(byte)',
    'fileType': '파일 형식', 'url': '파일 주소', 'readAt': '열람 일시',
    'memberNo': '회원번호', 'enrolledAt': '수강 등록일', 'purchasedAt': '구매 일시', 'totalSessions': '총 세션 수',
    'usedSessions': '사용 세션 수', 'usedAt': '사용 일시',
    'bio': '소개', 'photoUrl': '사진 주소', 'specialty': '전문 종목',
    'price': '가격(원)', 'pricingType': '이용 방식', 'ageGroup': '대상 연령대', 'capacity': '정원',
    'startTime': '시작 시각(HH:mm)', 'endTime': '종료 시각(HH:mm)',
    'cancelReason': '취소 사유', 'cancelledAt': '취소 일시',
    'amount': '결제 금액(원, 부가세 포함)', 'method': '결제 수단', 'mockApprovalNo': '모의 승인번호',
    'approvedAt': '승인 일시', 'refundedAt': '환불 일시',
    'currentCount': '현재 인원', 'lastUpdatedAt': '혼잡도 갱신 일시', 'recordedAt': '기록 일시', 'source': '기록 출처',
    'assetType': '자산 유형', 'acquiredAt': '취득일', 'acquisitionCost': '취득가액(원)', 'location': '보관 위치',
    'prefix': '코드 접두어', 'lastValue': '마지막 발급 번호', 'attemptedAt': '시도 일시',
    'action': '행위', 'entity': '대상 테이블', 'before': '변경 전(JSON)', 'after': '변경 후(JSON)',
}

DB_TYPE = {'String': 'text', 'Int': 'integer', 'Boolean': 'boolean', 'DateTime': 'timestamp(3)', 'Float': 'double precision',
           'Decimal': 'numeric', 'Json': 'jsonb', 'BigInt': 'bigint'}


def parse_schema():
    text = SCHEMA.read_text(encoding='utf-8')
    lines = text.split('\n')
    enums, models, domain = {}, {}, '공통'
    i = 0
    doc = []
    while i < len(lines):
        ln = lines[i].strip()
        m = re.match(r'// ── (?:\d+-\d+\. )?([^─(]+)', ln)
        if m and not ln.startswith('// ── D2'):
            domain = m.group(1).strip().rstrip('/').strip()
            domain = {'권한관리 / 조직': '권한관리', '게시판': '게시판', '공통: 변경 이력': '공통', '공통: 채번 시퀀스': '공통',
                      '공통: 회원 연동 실패 기록': '공통', '공통 enum': '공통', '예약 및 결제': '예약및결제',
                      '강사·프로그램 게시': '강사프로그램게시', '혼잡도 관리': '혼잡도관리', '자원문서관리': '자원문서관리'}.get(domain, domain)
            doc = []
        elif ln.startswith('//') and not ln.startswith('// ──'):
            doc.append(ln[2:].strip())
        elif ln.startswith('enum ') or ln.startswith('model '):
            kind, name = ln.split()[0], ln.split()[1]
            body = []
            i += 1
            while not lines[i].startswith('}'):
                body.append(lines[i])
                i += 1
            if kind == 'enum':
                vals = []
                for b in body:
                    b = b.strip()
                    if not b or b.startswith('//'):
                        continue
                    v, _, c = b.partition('//')
                    vals.append((v.strip(), c.strip()))
                enums[name] = {'domain': domain, 'values': vals, 'doc': ' '.join(doc)}
            else:
                models[name] = {'domain': domain, 'body': body, 'doc': ' '.join(doc)}
            doc = []
        elif ln == '':
            pass
        else:
            doc = []
        i += 1
    return enums, models


def parse_fields(models, enums):
    cols, rels, idx = [], [], []
    for mname, m in models.items():
        fk_of = {}
        body = m['body']
        # 관계 필드 먼저: @relation(fields: [...], references: [...], onDelete: X)
        for b in body:
            s = b.strip()
            r = re.search(r'@relation\((?:"[^"]*",\s*)?fields:\s*\[([^\]]+)\],\s*references:\s*\[([^\]]+)\](?:,\s*onDelete:\s*(\w+))?', s)
            if r:
                parent = s.split()[1].rstrip('?')
                for f in [x.strip() for x in r.group(1).split(',')]:
                    fk_of[f] = (parent, r.group(3) or '제한(기본)')
                rels.append([mname, r.group(1).replace(' ', ''), parent, r.group(3) or '제한(기본)', s.split()[0]])
        for b in body:
            s = b.strip()
            if not s or s.startswith('//'):
                continue
            if s.startswith('@@'):
                k = re.match(r'@@(\w+)\(\[([^\]]+)\]', s)
                if k:
                    idx.append([mname, {'index': '인덱스', 'unique': '유일(UNIQUE)', 'id': '기본키(복합)'}.get(k.group(1), k.group(1)),
                                k.group(2).replace(' ', ''), '', '스키마'])
                continue
            code, _, comment = s.partition('//')
            parts = code.split()
            if len(parts) < 2:
                continue
            fname, ftype = parts[0], parts[1]
            base = ftype.rstrip('?').rstrip('[]')
            if base in models:
                continue  # 객체 관계 필드는 컬럼이 아니다
            attrs = ' '.join(parts[2:])
            is_list = ftype.endswith('[]')
            nullable = ftype.endswith('?')
            if base in enums:
                dbt = f'enum {base}'
            else:
                dbt = DB_TYPE.get(base, base)
                if '@db.Date' in attrs:
                    dbt = 'date'
            if is_list:
                dbt += '[]'
            default = ''
            d = re.search(r'@default\((.*?)\)(?:\s|$)', attrs + ' ')
            if d:
                default = d.group(1)
            if '@updatedAt' in attrs:
                default = '(수정 시 자동)'
            pk = 'Y' if '@id' in attrs else ''
            uq = 'Y' if '@unique' in attrs else ''
            if uq:
                idx.append([mname, '유일(UNIQUE)', fname, '', '스키마'])
            fk = ''
            if fname in fk_of:
                fk = f'{fk_of[fname][0]}.id'
            desc = comment.strip() or GLOSSARY.get(fname, '')
            if fk and not comment.strip():
                desc = f'{TABLES.get(fk_of[fname][0], (fk_of[fname][0],))[0]} 참조' if fname not in GLOSSARY else GLOSSARY[fname]
            cols.append([m['domain'], mname, fname, ftype, dbt, pk, fk, 'Y' if nullable else 'N', default, uq, desc])
    return cols, rels, idx


def parse_sql():
    """CHECK·트리거·부분 유일 인덱스. 설명은 바로 앞 SQL 주석(-- ...)에서 가져온다."""
    checks, trigs, partial = [], [], []
    tok = re.compile(r'(?P<c>(?:^--[^\n]*\n)+)'
                     r'|(?P<chk>ALTER TABLE "(?P<ct>\w+)" ADD CONSTRAINT "(?P<cn>\w+)"\s+CHECK \((?P<ce>.*?)\);\s*$)'
                     r'|(?P<fn>CREATE (?:OR REPLACE )?FUNCTION "(?P<fname>\w+)")'
                     r'|(?P<tr>CREATE TRIGGER "(?P<tn>\w+)" (?P<tw>BEFORE|AFTER) (?P<te>.*?)\s+ON "(?P<tt>\w+)" FOR EACH ROW EXECUTE FUNCTION "(?P<tf>\w+)")'
                     r'|(?P<pu>CREATE UNIQUE INDEX "(?P<pn>\w+)" ON "(?P<pt>\w+)"\((?P<pc>[^)]*)\) WHERE (?P<pw>[^;]*);)',
                     re.M | re.S)
    for f in sorted(MIGR.glob('*/migration.sql')):
        sql = f.read_text(encoding='utf-8')
        mig = f.parent.name
        note, fn_note = '', {}
        for m in tok.finditer(sql):
            if m.group('c'):
                lines = [l.lstrip('-').strip() for l in m.group('c').strip().split('\n')]
                note = ' '.join(l for l in lines if l and not l.startswith('──'))
            elif m.group('chk'):
                checks.append([m.group('ct'), m.group('cn'), ' '.join(m.group('ce').split()).replace('"', ''), note, mig])
            elif m.group('fn'):
                fn_note[m.group('fname')] = note
            elif m.group('tr'):
                trigs.append([m.group('tt'), m.group('tn'), f"{m.group('tw')} {' '.join(m.group('te').split())}".replace('"', ''),
                              fn_note.get(m.group('tf'), note), mig])
            elif m.group('pu'):
                partial.append([m.group('pt'), '부분 유일(UNIQUE … WHERE)', m.group('pc').replace('"', ''),
                                ' '.join(m.group('pw').split()).replace('"', ''), mig])
    return checks, trigs, partial


def used_tables(models):
    """코드(apps/api/src, 테스트 제외)에서 실제로 읽고 쓰는지 — prisma 클라이언트 호출 또는 SQL의 "테이블명"."""
    src = '\n'.join(p.read_text(encoding='utf-8') for p in SRC.rglob('*.ts') if '.spec.' not in p.name and 'fixtures' not in p.parts)
    out = {}
    for name in models:
        camel = name[0].lower() + name[1:]
        out[name] = bool(re.search(rf'\.{camel}\.|"{name}"', src))
    return out


def main():
    enums, models = parse_schema()
    cols, rels, idx = parse_fields(models, enums)
    checks, trigs, partial = parse_sql()
    used = used_tables(models)
    commit = subprocess.run(['git', 'rev-parse', '--short', 'HEAD'], cwd=ROOT, capture_output=True, text=True).stdout.strip()

    wb = Workbook()
    # 1 표지
    ws = wb.active
    ws.title = '표지'
    r = X.title(ws, '스포이즘 ERP — 데이터 정의서', '테이블 정의서 · 코드 정의서 · 관계 · 제약')
    r = X.kv(ws, r, [
        ('프로젝트', '스포이즘 ERP — 아파트·오피스텔 커뮤니티 시설 위탁운영 관리 시스템'),
        ('문서', '데이터 정의서'),
        ('DBMS', 'PostgreSQL(Supabase, 서울 리전) · ORM: Prisma 5'),
        ('작성일', date.today().isoformat()),
        ('기준', f'{date.today().isoformat()} 저장소(dev 브랜치) — apps/api/prisma/schema.prisma, migrations/*.sql에서 자동 생성'),
        ('규모', f'테이블 {len(models)}개 · 컬럼 {len(cols)}개 · 코드(enum) {len(enums)}종 · CHECK {len(checks)}개 · 트리거 {len(trigs)}개'),
        ('미사용 테이블', ', '.join(f'{TABLES[n][0]}({n})' for n in models if not used[n]) + ' — 스키마에만 있고 코드가 읽거나 쓰지 않는다(요구사항 정의서 "미구현" 탭 참고)'),
    ])
    r += 1
    r = X.section(ws, r, '시트 구성', 2)
    r = X.kv(ws, r, [
        ('테이블 목록', '테이블별 도메인·논리명·설명·컬럼 수·사용 여부'),
        ('컬럼 정의', '전 테이블 컬럼 — 타입·PK·FK·NULL·기본값·UNIQUE·설명(필터로 테이블별 조회)'),
        ('관계', '외래키 — 자식 테이블, 참조 컬럼, 부모 테이블, 삭제 시 동작'),
        ('코드 정의', 'enum 코드값과 의미, 사용하는 컬럼'),
        ('인덱스·유일', '유일 제약, 부분 유일 인덱스(조건부), 조회 인덱스'),
        ('DB 규칙', 'CHECK 제약과 지점 일치 트리거 — 앱 검증과 이중으로 데이터 정합성을 지킨다'),
    ])
    X.widths(ws, [18, 110])

    # 2 테이블 목록
    ws = wb.create_sheet('테이블 목록')
    r = X.title(ws, '테이블 목록', '컬럼 수는 "컬럼 정의" 시트를 세는 수식이다')
    r = X.header(ws, r, ['No', '도메인', '테이블(물리명)', '논리명', '설명', '컬럼 수', '사용 여부'])
    start = r
    order = sorted(models, key=lambda n: (list(dict.fromkeys(m['domain'] for m in models.values())).index(models[n]['domain'])))
    data = []
    for k, n in enumerate(order, 1):
        data.append([k, models[n]['domain'], n, TABLES[n][0], TABLES[n][1], None, '사용' if used[n] else '미사용(설계만)'])
    r = X.rows(ws, r, data)
    for rr in range(start, r):
        ws.cell(row=rr, column=6, value=f"=COUNTIF('컬럼 정의'!$B:$B,C{rr})")
        c = ws.cell(row=rr, column=7)
        if c.value != '사용':
            c.fill = X.PatternFill('solid', fgColor=X.RED_SOFT)
    ws.cell(row=r, column=5, value='합계').font = X.font(True)
    ws.cell(row=r, column=6, value=f'=SUM(F{start}:F{r-1})').font = X.font(True)
    X.widths(ws, [6, 16, 26, 18, 70, 10, 16])
    ws.freeze_panes = ws.cell(row=start, column=1)

    # 3 컬럼 정의
    ws = wb.create_sheet('컬럼 정의')
    r = X.title(ws, '컬럼 정의', '헤더의 필터로 테이블·도메인별로 볼 수 있다. 같은 테이블은 같은 줄무늬')
    hdr = ['도메인', '테이블', '컬럼', 'Prisma 타입', 'DB 타입', 'PK', 'FK(참조)', 'NULL 허용', '기본값', 'UNIQUE', '설명']
    r = X.header(ws, r, hdr)
    start = r
    cols.sort(key=lambda c: (order.index(c[1]),))
    r = X.rows(ws, r, cols, band_key=1)
    ws.auto_filter.ref = f'A{start-1}:K{r-1}'
    X.widths(ws, [14, 24, 24, 16, 22, 5, 22, 9, 22, 8, 60])
    ws.freeze_panes = ws.cell(row=start, column=4)

    # 4 관계
    ws = wb.create_sheet('관계')
    r = X.title(ws, '관계(외래키)', '삭제 시 동작: Cascade = 부모 삭제 시 함께 삭제, SetNull = NULL로, 제한(기본) = 자식이 있으면 부모 삭제 불가')
    r = X.header(ws, r, ['자식 테이블', 'FK 컬럼', '부모 테이블', '삭제 시 동작', '관계 필드명'])
    rels.sort(key=lambda x: order.index(x[0]))
    r = X.rows(ws, r, rels, band_key=0)
    X.widths(ws, [26, 26, 26, 16, 24])

    # 5 코드 정의
    ws = wb.create_sheet('코드 정의')
    r = X.title(ws, '코드 정의(enum)', '코드값은 DB에 영문으로 저장되고 화면에는 의미(한글)로 표시된다')
    r = X.header(ws, r, ['코드 그룹', '도메인', '코드값', '의미', '사용 컬럼'])
    usage = {}
    for c in cols:
        if c[4].startswith('enum '):
            usage.setdefault(c[4][5:].rstrip('[]'), []).append(f'{c[1]}.{c[2]}')
    data = []
    for e, v in enums.items():
        for val, meaning in v['values']:
            data.append([e, v['domain'], val, meaning, ', '.join(usage.get(e, []))])
    r = X.rows(ws, r, data, band_key=0)
    X.widths(ws, [24, 14, 22, 50, 50])

    # 6 인덱스·유일
    ws = wb.create_sheet('인덱스·유일')
    r = X.title(ws, '인덱스·유일 제약', '부분 유일 인덱스는 Prisma 스키마로 표현할 수 없어 마이그레이션 SQL에만 있다(조건을 만족하는 행끼리만 유일)')
    r = X.header(ws, r, ['테이블', '종류', '컬럼', '조건', '출처'])
    allidx = sorted(idx + partial, key=lambda x: (order.index(x[0]) if x[0] in order else 99, x[1]))
    r = X.rows(ws, r, allidx, band_key=0)
    X.widths(ws, [26, 24, 40, 50, 40])

    # 7 DB 규칙
    ws = wb.create_sheet('DB 규칙')
    r = X.title(ws, 'DB 규칙 — CHECK 제약 · 트리거', '앱이 먼저 검증해 400/409를 내고, DB가 한 번 더 막는다(D28). 규칙 목록·원칙: architecture/data-integrity.md')
    r = X.section(ws, r, f'CHECK 제약 {len(checks)}개', 4)
    r = X.header(ws, r, ['테이블', '제약명', '조건', '근거·설명', '마이그레이션'])
    r = X.rows(ws, r, checks, band_key=0)
    r += 1
    r = X.section(ws, r, f'지점 일치 트리거 {len(trigs)}개 — 연결된 데이터는 같은 지점이어야 한다', 4)
    r = X.header(ws, r, ['테이블', '트리거명', '시점·대상', '설명', '마이그레이션'])
    r = X.rows(ws, r, trigs, band_key=0)
    X.widths(ws, [20, 34, 60, 70, 36])

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'{OUT.relative_to(ROOT)} — 테이블 {len(models)}, 컬럼 {len(cols)}, enum {len(enums)}, 관계 {len(rels)}, '
          f'인덱스 {len(allidx)}, CHECK {len(checks)}, 트리거 {len(trigs)}, 미사용 {sum(not v for v in used.values())}')


if __name__ == '__main__':
    main()
