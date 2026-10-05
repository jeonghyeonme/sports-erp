"""요구사항 정의서(엑셀) 생성 — 원천은 저장소다.

    python3 scripts/deliverables/requirements.py   → docs/deliverables/요구사항정의서.xlsx

- RFP 추적·미구현·범위 제외: docs/reference/요구사항추적표.md §1·§2-3·§2-4
- 도메인별 원문 요구(기능·품질)·불변규칙·설계 결정·범위 제외: docs/domains/<도메인>.md
- 기능 명세: apps/api/src의 실제 컨트롤러 라우트(경로·권한) + 관리자 웹이 실제로 부르는 경로(화면 유무)
  — 설명(ROUTE_DESC)만 사람이 쓴다. 라우트가 추가되면 설명이 비어 경고가 뜬다.
- 자동 테스트: apps/api/test/*.spec.ts 파일명
"""
import re
import subprocess
from datetime import date
from pathlib import Path

from openpyxl import Workbook

import xlsx_style as X

ROOT = Path(__file__).resolve().parents[2]
DOMAINS = ROOT / 'docs/domains'
TRACE = ROOT / 'docs/reference/요구사항추적표.md'
API = ROOT / 'apps/api/src'
WEB = ROOT / 'apps/admin-web/src'
TESTS = ROOT / 'apps/api/test'
OUT = ROOT / 'docs/deliverables/요구사항정의서.xlsx'

# 도메인 — (시트 이름, 문서 파일, ID 접두어, 컨트롤러 경로 접두어)
DOMAIN_LIST = [
    ('권한관리', '권한관리.md', 'AUTH', ['/auth', '/permissions']),
    ('인사정보관리', '인사정보관리.md', 'STF', ['/staff']),
    ('근태관리', '근태관리.md', 'ATT', ['/attendance', '/leave-requests', '/leave-balance', '/work-logs']),
    ('게시판', '게시판.md', 'BRD', ['/posts']),
    ('회원관리', '회원관리.md', 'MEM', ['/members']),
    ('예약및결제', '예약및결제.md', 'RSV', ['/reservations', '/payments']),
    ('강사프로그램게시', '강사프로그램게시.md', 'PRG', ['/programs', '/instructors', '/branches/:branchId/programs']),
    ('혼잡도관리', '혼잡도관리.md', 'FAC', ['/facilities']),
    ('자원문서관리', '자원문서관리.md', 'RES', ['/assets', '/documents']),
]
COMMON_PREFIX = ['/branches', '/health']

ROUTE_DESC = {
    'POST /auth/login': '이메일·비밀번호 로그인 — 액세스·리프레시 토큰 발급, 비활성 계정 차단',
    'POST /auth/refresh': '리프레시 토큰으로 재발급(1회용 교체)',
    'POST /auth/logout': '로그아웃 — 리프레시 토큰 폐기',
    'GET /auth/me': '내 계정 정보',
    'PATCH /auth/password': '비밀번호 변경',
    'GET /permissions/staff': '직원 계정과 역할 목록(본사)',
    'PATCH /permissions/staff/:staffId/role': '직원 역할 전환(지점 관리자 ↔ 직원, 본사)',
    'GET /staff': '직원 목록(본사 전체, 지점 관리자는 자기 지점)',
    'GET /staff/me': '내 직원 정보',
    'GET /staff/:id': '직원 상세',
    'POST /staff': '직원 채용 — 계정·직원·첫 파견을 함께 생성(본사만)',
    'PATCH /staff/:id': '직원 정보 수정',
    'PATCH /staff/:id/resign': '퇴사 처리 — 계정 비활성화, 파견 종료, 인사서류 보존기한 재계산',
    'POST /staff/:id/assignments': '파견 발령(재배치) — 옛 지점 담당 회원 해제(본사만)',
    'GET /staff/:id/assignments': '파견 이력',
    'POST /attendance/check-in': '출근 체크 — 하루 1회, 지각 자동 판정',
    'POST /attendance/check-out': '퇴근 체크',
    'GET /attendance': '근태 기록 조회(범위별)',
    'GET /attendance/summary': '지점 월별 근태 요약',
    'GET /attendance/absence-preview': '결근 예정 미리보기(저장 안 함)',
    'POST /attendance/absence-confirm': '결근 확정(지점 관리자의 명시적 확정)',
    'POST /leave-requests': '휴가 신청',
    'GET /leave-requests': '휴가 신청 목록',
    'PATCH /leave-requests/:id/approve': '휴가 승인 — 승인 시점에 연차 차감',
    'PATCH /leave-requests/:id/reject': '휴가 반려',
    'GET /leave-balance/:staffId': '연차 잔여 조회',
    'POST /work-logs': '업무일지 작성(하루 1건)',
    'GET /work-logs': '업무일지 조회',
    'GET /posts': '게시글 목록 — 범위별 권한, 회원 공개 필터, 페이지네이션',
    'GET /posts/:id': '게시글 상세 — 조회수 증가',
    'POST /posts': '게시글 작성 — 계약 종료 지점 차단',
    'PATCH /posts/:id': '게시글 수정(작성자)',
    'DELETE /posts/:id': '게시글 삭제(소프트, 작성자 또는 본사)',
    'POST /members/register': '회원 앱 자체 가입(공개) — 회원 앱 범위 제외로 화면 없음',
    'POST /members/link': '오프라인 회원 ↔ 앱 계정 연동(공개, 시도 제한) — 회원 앱 범위 제외로 화면 없음',
    'GET /members': '회원 목록(지점 범위)',
    'GET /members/:id': '회원 상세 + 요약(수강·PT·예약)',
    'POST /members': '회원 등록 — 미성년 보호자 동의 확인, 계약 종료 지점 차단',
    'PATCH /members/:id': '회원 정보 수정',
    'PATCH /members/:id/status': '회원 상태 전환(활성·휴면·탈퇴)',
    'GET /members/:id/enrollments': '수강 내역',
    'POST /members/:id/enrollments': '수강 등록',
    'GET /members/:id/pt-sessions': 'PT 세션권 목록',
    'POST /members/:id/pt-sessions': 'PT 세션권 등록',
    'POST /members/:id/pt-sessions/:sessionId/use': 'PT 1회 차감',
    'POST /reservations': '회차 예약 — 회차 잠금, 지점·계약·정원·중복 확인',
    'GET /reservations': '예약 목록(범위별)',
    'PATCH /reservations/:id/cancel': '예약 취소 — 마감 전이면 환불',
    'PATCH /reservations/:id/check-in': '현장 체크인(완료 처리)',
    'POST /payments/:reservationId/mock-pay': '모의 결제 승인 → 예약 확정(대기일 때 한 번만)',
    'GET /payments': '결제 내역(매출)',
    'GET /programs': '프로그램 목록',
    'POST /programs': '프로그램 등록',
    'PATCH /programs/:id': '프로그램 수정',
    'PATCH /programs/:id/status': '프로그램 상태 전이(준비중·진행중·휴강·종료) — 영향받는 예약 반환',
    'DELETE /programs/:id': '프로그램 종료(소프트)',
    'GET /programs/:id/slots': '회차 목록·잔여 좌석',
    'POST /programs/:id/slots': '회차 등록',
    'GET /instructors': '강사 목록',
    'POST /instructors': '강사 등록',
    'PATCH /instructors/:id': '강사 수정',
    'DELETE /instructors/:id': '강사 비활성화',
    'GET /branches/:branchId/programs/summary': '지점별 프로그램 상태 요약',
    'GET /facilities': '시설 목록 + 현재 혼잡도',
    'POST /facilities': '시설 등록',
    'PATCH /facilities/:id': '시설 수정·비활성화',
    'POST /facilities/:id/congestion/manual': '혼잡도 수동 보정(인원 입력 → 5단계 자동 계산)',
    'GET /assets': '자산·비품 목록',
    'POST /assets': '자산 등록 — 자산코드 자동 채번, 계약 종료 지점 차단',
    'PATCH /assets/:id': '자산 정보 수정',
    'PATCH /assets/:id/status': '자산 상태 전이(사용·수리·폐기 등)',
    'GET /documents': '문서 목록',
    'GET /documents/retention-alerts': '보존기한 임박 문서(본사)',
    'GET /documents/:id': '문서 상세',
    'POST /documents': '문서 등록 — 계약서 보존기한 필수, 계약 종료 지점 차단',
    'DELETE /documents/:id': '문서 삭제(소프트, 본사)',
    'GET /branches': '지점 목록(계약 정보 포함)',
    'GET /health': '헬스체크(DB 미경유)',
}
ROLE_KO = {'SUPER_ADMIN': '본사', 'BRANCH_ADMIN': '지점', 'STAFF': '직원', 'MEMBER': '회원'}


def read(p):
    return Path(p).read_text(encoding='utf-8')


def md_table(lines):
    rows = []
    for l in lines:
        l = l.strip()
        if not l.startswith('|') or re.match(r'^\|[-\s|:]+\|$', l):
            continue
        rows.append([c.strip() for c in l.strip('|').split('|')])
    return rows


def clean(t):
    t = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', t)       # 링크
    t = re.sub(r'\*\*|`|~~', '', t)                        # 강조·코드·취소선
    return t.strip()


def section_lines(text, start_pat, end_pat):
    out, on = [], False
    for l in text.split('\n'):
        if re.match(start_pat, l):
            on = True
            continue
        if on and re.match(end_pat, l):
            break
        if on:
            out.append(l)
    return out


# ── 코드 실측 ──────────────────────────────────────────────────────────────
def api_routes():
    routes = []
    for f in sorted(API.rglob('*.controller.ts')):
        s = read(f)
        pre = re.search(r"@Controller\('?([^)']*)'?\)", s).group(1)
        cls = re.search(r"((?:@[^\n]*\n)*)export class", s)
        croles = re.findall(r"@Roles\(([^)]*)\)", cls.group(1)) if cls else []
        lines = s.split('\n')
        for i, l in enumerate(lines):
            m = re.match(r"\s*@(Get|Post|Patch|Put|Delete)\((?:'([^']*)')?\)", l)
            if not m:
                continue
            # 메서드 위아래 데코레이터 묶음
            block, j = [l], i - 1
            while j >= 0 and lines[j].strip().startswith('@'):
                block.append(lines[j]); j -= 1
            j = i + 1
            while j < len(lines) and lines[j].strip().startswith('@'):
                block.append(lines[j]); j += 1
            b = '\n'.join(block)
            roles = re.findall(r"@Roles\(([^)]*)\)", b) or croles
            path = '/' + '/'.join(x for x in [pre, m.group(2) or ''] if x)
            if '@Public' in b:
                who = '공개'
            elif roles:
                who = ', '.join(ROLE_KO.get(r.strip(" '"), r) for r in roles[0].split(','))
            else:
                who = '로그인 사용자'
            routes.append((m.group(1).upper(), path, who))
    return routes


def web_paths():
    """관리자 웹이 부르는 API 경로. 쓰기(POST/PATCH/DELETE)는 api.<method>(경로) 호출로, 조회(GET)는
    경로 문자열이 코드에 나오면(useApiList 같은 조회 헬퍼가 경로를 인자로 받는다) 화면이 있다고 본다."""
    s = '\n'.join(read(p) for p in list(WEB.rglob('*.tsx')) + list(WEB.rglob('*.ts')))
    norm = lambda x: re.sub(r'\$\{[^}]+\}', ':p', x).split('?')[0].rstrip('/') or '/'
    out = set()
    for m in re.finditer(r"api\.(get|post|patch|put|delete)(?:<[^(]*>)?\(\s*[`']([^`']+)[`']", s):
        out.add((m.group(1).upper(), norm(m.group(2))))
    for m in re.finditer(r"[`'](/[a-z][^`'\s]*)[`']", s):
        out.add(('GET', norm(m.group(1))))
    return out


def has_screen(method, path, web):
    seg = path.strip('/').split('/')
    for wm, wp in web:
        if wm != method:
            continue
        ws_ = wp.strip('/').split('/')
        if len(ws_) == len(seg) and all(a == b or a == ':p' or b.startswith(':') for a, b in zip(ws_, seg)):
            return True
    return False


def tests():
    return sorted(p.stem.replace('.spec', '') for p in TESTS.glob('*.spec.ts'))


# ── 문서 파싱 ──────────────────────────────────────────────────────────────
def parse_domain(fname):
    t = read(DOMAINS / fname)
    sec1 = section_lines(t, r'^## 1\.', r'^## 2\.')
    reqs, inv, kind = [], [], None
    for l in sec1:
        if l.startswith('### '):
            kind = l[4:].split(' ')[0]
            continue
        if kind in ('기능', '품질'):
            r = md_table([l])
            if r and r[0][0] != 'ID':
                reqs.append([r[0][0], kind, clean(r[0][1])])
        elif kind == '제약' and l.startswith('- '):
            reqs.append(['', '제약', clean(l[2:])])
        elif kind == '불변규칙':
            m = re.match(r'^\d+\.\s+(.*)', l)
            if m:
                inv.append(clean(m.group(1)))
    card = section_lines(t, r'^## 요약 카드', r'^## 0\.')
    adr = [[clean(c) for c in r[:3]] for r in md_table(card) if r and r[0].startswith('ADR-')]
    test_line = next((clean(l.split('**테스트**:', 1)[1]) for l in card if '**테스트**' in l), '')
    gaps = [clean(l.split(':', 1)[1]) for l in card if re.match(r'^- \*\*(미구현|범위 제외)', l)]
    stories = [clean(l[2:]) for l in section_lines(t, r'^## 0\.', r'^## 1\.') if l.startswith('- ')]
    return {'reqs': reqs, 'inv': inv, 'adr': adr, 'tests': test_line, 'gaps': gaps, 'stories': stories}


def parse_trace():
    t = read(TRACE)
    s11 = md_table(section_lines(t, r'^### 1-1\.', r'^### 1-2\.'))[1:]
    s12 = md_table(section_lines(t, r'^### 1-2\.', r'^### 1-3\.'))[1:]
    s13 = md_table(section_lines(t, r'^### 1-3\.', r'^---'))[1:]
    s23 = md_table(section_lines(t, r'^### 2-3\.', r'^### 2-4\.'))[1:]
    s24 = md_table(section_lines(t, r'^### 2-4\.', r'^---'))[1:]
    api_only = next((clean(l.split('**:', 1)[1]) for l in section_lines(t, r'^### 2-3\.', r'^### 2-4\.') if l.startswith('**API는 있고')), '')
    design_only = next((clean(l.split('**:', 1)[1]) for l in section_lines(t, r'^### 2-3\.', r'^### 2-4\.') if l.startswith('**설계는 있지만')), '')
    member_app = next((clean(l) for l in section_lines(t, r'^### 2-4\.', r'^---') if l.startswith('**회원 앱 전체')), '')
    return s11, s12, s13, s23, s24, api_only, design_only, member_app


# ── 작성 ──────────────────────────────────────────────────────────────────
def main():
    routes = api_routes()
    web = web_paths()
    commit = subprocess.run(['git', 'rev-parse', '--short', 'HEAD'], cwd=ROOT, capture_output=True, text=True).stdout.strip()
    s11, s12, s13, s23, s24, api_only, design_only, member_app = parse_trace()
    missing_desc = [f'{m} {p}' for m, p, _ in routes if f'{m} {p}' not in ROUTE_DESC]
    all_rows = []  # 총괄

    wb = Workbook()
    ws = wb.active
    ws.title = '표지'

    # 도메인 시트
    used = set()
    domain_sheets = []
    for name, fname, code, prefixes in DOMAIN_LIST:
        d = parse_domain(fname)
        sh = wb.create_sheet(name)
        domain_sheets.append(name)
        r = X.title(sh, f'{name} — 요구사항', f'원천: docs/domains/{fname} · 기능 명세는 코드 실측(커밋 {commit})')
        if d['stories']:
            r = X.section(sh, r, '사용자 스토리', 6)
            for st in d['stories']:
                sh.cell(row=r, column=1, value='•').font = X.font()
                c = sh.cell(row=r, column=2, value=st); c.font = X.font()
                sh.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
                r += 1
            r += 1
        # 원문 요구
        r = X.section(sh, r, '원문 요구(RFP 기반 재분류)', 6)
        r = X.header(sh, r, ['요구사항 ID', '구분', '내용', '', '', ''])
        data = []
        for k, (rid, kind, text) in enumerate(d['reqs'], 1):
            rid = rid or f'{code}-C{k:02d}'
            data.append([rid, kind, text, '', '', ''])
            all_rows.append([rid, name, f'원문 {kind}', text, '-', '-', ''])
        r = X.rows(sh, r, data)
        for rr in range(r - len(data), r):
            sh.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=6)
        r += 1
        # 불변규칙
        r = X.section(sh, r, '불변규칙(협의 불가)', 6)
        r = X.header(sh, r, ['요구사항 ID', '규칙', '', '', '', '검증'])
        data = []
        for k, t in enumerate(d['inv'], 1):
            rid = f'{code}-INV-{k:02d}'
            data.append([rid, t, '', '', '', '자동 테스트' if d['tests'] else '검증되지 않음'])
            all_rows.append([rid, name, '불변규칙', t, '구현', '-', d['tests']])
        r = X.rows(sh, r, data)
        for rr in range(r - len(data), r):
            sh.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
        sh.cell(row=r, column=1, value='관련 테스트').font = X.font(True)
        c = sh.cell(row=r, column=2, value=d['tests'] or '없음'); c.font = X.font()
        sh.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        r += 2
        # 기능 명세
        r = X.section(sh, r, '기능 명세(실제 API 기준)', 6)
        r = X.header(sh, r, ['요구사항 ID', '기능', 'Method', 'Endpoint(/api/v1)', '권한', '상태'])
        data = []
        mine = [x for x in routes if any(x[1] == p or x[1].startswith(p + '/') for p in prefixes)]
        mine.sort(key=lambda x: (next(i for i, p in enumerate(prefixes) if x[1] == p or x[1].startswith(p + '/')), x[1], x[0]))
        for k, (m, p, who) in enumerate(mine, 1):
            used.add((m, p))
            rid = f'{code}-FN-{k:02d}'
            desc = ROUTE_DESC.get(f'{m} {p}', '(설명 없음)')
            if has_screen(m, p, web):
                st = '구현(API+화면)'
            elif '/register' in p or '/link' in p:
                st = 'API만(회원 앱 범위 제외)'
            elif p.startswith('/auth/') and p != '/auth/login':
                st = '구현(API, 화면 불필요)' if p in ('/auth/refresh', '/auth/logout', '/auth/me') else 'API만(화면 없음)'
            else:
                st = 'API만(화면 없음)'
            data.append([rid, desc, m, p, who, st])
            all_rows.append([rid, name, '기능', desc, '구현', '있음' if '화면' in st and 'API만' not in st else ('불필요' if '불필요' in st else '없음'), f'{m} {p}'])
        r = X.rows(sh, r, data, status_col=5)
        r += 1
        # 설계 결정
        r = X.section(sh, r, '설계 결정(ADR)', 6)
        r = X.header(sh, r, ['ADR', '결정', '', '', '', '상태'])
        data = [[a[0], a[1], '', '', '', a[2]] for a in d['adr']]
        r = X.rows(sh, r, data, status_col=5)
        for rr in range(r - len(data), r):
            sh.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
        if d['gaps']:
            r += 1
            r = X.section(sh, r, '미구현·범위 제외(도메인 문서 기준)', 6)
            for g in d['gaps']:
                c = sh.cell(row=r, column=1, value=g); c.font = X.font()
                c.alignment = X.Alignment(wrap_text=True, vertical='top')
                sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
                sh.row_dimensions[r].height = 30
                r += 1
        X.widths(sh, [16, 46, 9, 40, 22, 26])
        sh.sheet_view.showGridLines = False

    # 공통 라우트
    common = [x for x in routes if (x[0], x[1]) not in used]
    for k, (m, p, who) in enumerate(common, 1):
        all_rows.append([f'COM-FN-{k:02d}', '공통', '기능', ROUTE_DESC.get(f'{m} {p}', '(설명 없음)'), '구현',
                         '불필요' if p == '/health' else ('있음' if has_screen(m, p, web) else '없음'), f'{m} {p}'])

    # 미구현을 총괄에 추가
    for k, row in enumerate(s23, 1):
        all_rows.append([f'GAP-{k:02d}', 'RFP 공통', '미구현', f'{clean(row[1])} — {clean(row[2])}', '미구현', '-', clean(row[0])])

    # 총괄
    sh = wb.create_sheet('요구사항 총괄', 1)
    r = X.title(sh, '요구사항 총괄', '헤더 필터로 도메인·구분·상태별로 볼 수 있다. 위 집계는 아래 표를 세는 수식이다')
    stat_row = r
    labels = [('기능(API) 수', '기능', None), ('  └ 화면 있음', '기능', '있음'), ('  └ 화면 없음(API만)', '기능', '없음'),
              ('불변규칙 수', '불변규칙', None), ('RFP 미구현 항목', '미구현', None)]
    r = X.header(sh, r, ['집계', '건수'])
    first_data_row = r + len(labels) + 2
    last_guess = first_data_row + len(all_rows)
    for lab, kind, scr in labels:
        sh.cell(row=r, column=1, value=lab).font = X.font(True)
        if scr:
            f = f'=COUNTIFS($C${first_data_row}:$C${last_guess},"{kind}",$F${first_data_row}:$F${last_guess},"{scr}")'
        else:
            f = f'=COUNTIF($C${first_data_row}:$C${last_guess},"{kind}")'
        c = sh.cell(row=r, column=2, value=f); c.font = X.font(); c.border = X.BORDER
        r += 1
    r += 1
    r = X.header(sh, r, ['요구사항 ID', '도메인', '구분', '내용', '상태', '화면', '근거·검증'])
    assert r == first_data_row, (r, first_data_row)
    r = X.rows(sh, r, all_rows, status_col=4, band_key=1)
    sh.auto_filter.ref = f'A{first_data_row-1}:G{r-1}'
    sh.freeze_panes = sh.cell(row=first_data_row, column=1)
    X.widths(sh, [16, 16, 12, 70, 10, 8, 50])
    sh.sheet_view.showGridLines = False

    # RFP 추적
    sh = wb.create_sheet('RFP 추적', 2)
    r = X.title(sh, 'RFP 추적표', '제안요청서(2022.06) 원문 요구 → 도메인 → 조정 내용과 근거. 원천: docs/reference/요구사항추적표.md §1')
    r = X.section(sh, r, '기능·성능 요구(원문 순번, 9번은 원문 결번)', 6)
    r = X.header(sh, r, ['#', '구분', '요구사항명', '원문 요약', '도메인', '조정 요약', '결정 근거'])
    r = X.rows(sh, r, [[clean(c) for c in (row[:3] + [row[3], row[4], row[5], row[6]])] for row in s11])
    r += 1
    r = X.section(sh, r, '시스템·데이터·인터페이스·보안·사업지원·제약 요구', 6)
    r = X.header(sh, r, ['#', '구분', '요구사항명', '원문 요약', '대응·조정'])
    r = X.rows(sh, r, [[clean(c) for c in row[:5]] for row in s12])
    r += 1
    r = X.section(sh, r, '과업1·과업2·산출물', 6)
    r = X.header(sh, r, ['원본 항목', '위치', '대응 문서', '비고'])
    r = X.rows(sh, r, [[clean(c) for c in row[:4]] for row in s13])
    X.widths(sh, [10, 12, 22, 50, 24, 60, 16])
    sh.sheet_view.showGridLines = False

    # 미구현
    sh = wb.create_sheet('미구현')
    r = X.title(sh, '미구현 — RFP가 요구하는데 아직 없는 것', '의도적 범위 제외(다음 시트)와 구분한다. 원천: 요구사항추적표 §2-3(코드 실측)')
    r = X.header(sh, r, ['ID', 'RFP', '항목', '지금 상태(실측)', '상태'])
    r = X.rows(sh, r, [[f'GAP-{k:02d}', clean(row[0]), clean(row[1]), clean(row[2]), '미구현'] for k, row in enumerate(s23, 1)], status_col=4)
    r += 1
    r = X.section(sh, r, 'API는 있고 화면이 없는 것', 5)
    c = sh.cell(row=r, column=1, value=api_only); c.font = X.font(); c.alignment = X.Alignment(wrap_text=True, vertical='top')
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); sh.row_dimensions[r].height = 36
    r += 2
    r = X.section(sh, r, '설계는 있지만 미구현인 도메인 규칙', 5)
    c = sh.cell(row=r, column=1, value=design_only); c.font = X.font(); c.alignment = X.Alignment(wrap_text=True, vertical='top')
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5); sh.row_dimensions[r].height = 36
    X.widths(sh, [10, 26, 30, 70, 10])
    sh.sheet_view.showGridLines = False

    # 범위 제외
    sh = wb.create_sheet('범위 제외')
    r = X.title(sh, '범위 제외 — 의도적으로 만들지 않은 것', '원천: 요구사항추적표 §2-4, 각 도메인 문서 부록 A-8')
    r = X.header(sh, r, ['도메인', '범위 제외 항목', '향후 확장'])
    r = X.rows(sh, r, [[clean(c) for c in row[:3]] for row in s24], status_col=None)
    r += 1
    r = X.section(sh, r, '회원 앱', 3)
    c = sh.cell(row=r, column=1, value=member_app + ' 보류 이유: 회원 앱을 받으면 하루 요청이 약 19.6만 건(계산값)으로 엣지 무료 한도 10만 건을 넘는다. '
                                                  '유료 전환이나 인프라를 더 구성해 연결해야 해 개발 범위가 과하게 넓어지므로 보류했다(D37 §5, log/057).')
    c.font = X.font(); c.alignment = X.Alignment(wrap_text=True, vertical='top')
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); sh.row_dimensions[r].height = 90
    X.widths(sh, [18, 80, 50])
    sh.sheet_view.showGridLines = False

    # 표지
    ws = wb['표지']
    r = X.title(ws, '스포이즘 ERP — 요구사항 정의서', '요구사항 총괄 · RFP 추적 · 도메인별 요구사항 · 미구현 · 범위 제외')
    n_fn = sum(1 for x in all_rows if x[2] == '기능')
    r = X.kv(ws, r, [
        ('프로젝트', '스포이즘 ERP — 아파트·오피스텔 커뮤니티 시설 위탁운영 관리 시스템'),
        ('문서', '요구사항 정의서'),
        ('작성일', date.today().isoformat()),
        ('기준', f'저장소 커밋 {commit} — 도메인 문서·요구사항추적표·실제 API 라우트에서 자동 생성'),
        ('규모', f'도메인 {len(DOMAIN_LIST)}개 · 기능(API) {n_fn}개 · 불변규칙 {sum(1 for x in all_rows if x[2]=="불변규칙")}개 · '
               f'RFP 미구현 {len(s23)}건 · 자동 테스트 파일 {len(tests())}개'),
        ('범위', '관리자 웹(본사·지점 관리자·직원) + REST API. 회원 앱은 범위 제외(회원 기능은 API로 구현, 회원 계정으로 관리자 웹에서 사용)'),
    ])
    r += 1
    r = X.section(ws, r, '요구사항 ID 규칙', 2)
    r = X.kv(ws, r, [
        ('F·Q + 번호', 'RFP 원문을 재분류한 기능(F)·품질(Q) 요구 — 도메인 문서 §1의 ID 그대로'),
        ('<도메인>-INV-nn', '불변규칙 — 어떤 경우에도 깨지면 안 되는 규칙'),
        ('<도메인>-FN-nn', '기능 명세 — 실제 API 하나에 하나'),
        ('GAP-nn', 'RFP가 요구하는데 아직 구현하지 않은 항목'),
        ('도메인 약어', ', '.join(f'{c}={n}' for n, _, c, _ in DOMAIN_LIST) + ', COM=공통'),
    ])
    r += 1
    r = X.section(ws, r, '상태 표기', 2)
    r = X.kv(ws, r, [
        ('구현(API+화면)', 'API와 관리자 웹 화면이 모두 있다'),
        ('API만', 'API는 있고 화면이 없다(사유를 괄호에 적음)'),
        ('미구현', 'RFP 요구인데 아직 없다'),
        ('범위 제외', '근거를 두고 의도적으로 만들지 않았다'),
        ('검증', '자동 테스트는 API(jest)만 있다. 관리자 웹은 lint·빌드(타입체크)까지만 검증된다'),
    ])
    X.widths(ws, [20, 110])
    ws.sheet_view.showGridLines = False

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f'{OUT.relative_to(ROOT)} — 라우트 {len(routes)}(공통 {len(common)}), 총괄 {len(all_rows)}행, 시트 {wb.sheetnames}')
    if missing_desc:
        print('경고: 설명 없는 라우트', missing_desc)


if __name__ == '__main__':
    main()
