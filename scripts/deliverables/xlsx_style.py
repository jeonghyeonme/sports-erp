"""산출물 엑셀 공통 서식 — 데이터 정의서·요구사항 정의서가 같은 모양을 쓰게 한다."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
NAVY = '1E3A5F'
NAVY_SOFT = 'E7ECF1'
GRAY_SOFT = 'F2F4F3'
GREEN_SOFT = 'E3F3EB'
AMBER_SOFT = 'FBEEDD'
RED_SOFT = 'F8E6E4'
THIN = Side(style='thin', color='CBD0CD')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILL = {
    '구현': GREEN_SOFT, '적용': GREEN_SOFT,
    '미구현': RED_SOFT, '설계만': RED_SOFT,
    'API만': AMBER_SOFT, '부분': AMBER_SOFT, '측정 전': AMBER_SOFT,
    '범위 제외': GRAY_SOFT,
}


def font(bold=False, size=10, color='14191C'):
    return Font(name=FONT, bold=bold, size=size, color=color)


def title(ws, text, sub=None):
    ws['A1'] = text
    ws['A1'].font = font(True, 16, NAVY)
    if sub:
        ws['A2'] = sub
        ws['A2'].font = font(False, 10, '4B5563')
    return 4  # 다음에 쓸 행


def section(ws, row, text, ncols):
    c = ws.cell(row=row, column=1, value=text)
    c.font = font(True, 12, NAVY)
    return row + 1


def header(ws, row, names):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = font(True, 10, 'FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
    return row + 1


def rows(ws, row, data, status_col=None, band_key=None):
    """data: list[list]. status_col: 0-based 열 번호 — 값에 따라 칸 색을 칠한다. band_key: 같은 값끼리 줄무늬."""
    last, band = object(), False
    for r in data:
        if band_key is not None and r[band_key] != last:
            band, last = not band, r[band_key]
        for i, v in enumerate(r, 1):
            c = ws.cell(row=row, column=i, value=v)
            c.font = font()
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = BORDER
            if band_key is not None and band:
                c.fill = PatternFill('solid', fgColor=GRAY_SOFT)
        if status_col is not None:
            v = r[status_col]
            for k, color in STATUS_FILL.items():
                if isinstance(v, str) and v.startswith(k):
                    ws.cell(row=row, column=status_col + 1).fill = PatternFill('solid', fgColor=color)
                    break
        row += 1
    return row


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def kv(ws, row, pairs):
    """표지용 키-값 표."""
    for k, v in pairs:
        a = ws.cell(row=row, column=1, value=k)
        a.font = font(True)
        a.fill = PatternFill('solid', fgColor=NAVY_SOFT)
        a.border = BORDER
        b = ws.cell(row=row, column=2, value=v)
        b.font = font()
        b.alignment = Alignment(wrap_text=True, vertical='top')
        b.border = BORDER
        row += 1
    return row
